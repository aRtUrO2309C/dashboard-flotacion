"""Ejecutar: python -m streamlit run Dasboard_Leyes_Relaves.py --server.port 8502
Instalar en el entorno: python -m pip install streamlit pandas plotly openpyxl
Lee el Excel en cada interacción. No modifica el libro ni sus macros.
"""
from pathlib import Path
from datetime import datetime, date
import math
import re as regex
import openpyxl
import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import streamlit as st

CARPETA = Path(__file__).resolve().parent
NOMBRE_EXCEL = "Analisis Flotacion por Guardia.xlsm"
HOJA = "Analisis Sulfuros por Guardia"
FECHA_INICIO = pd.Timestamp("2026-06-23")
MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
COLORES = {"Día": "#1645df", "Noche": "#f56b18"}
COLORES_SUPERVISOR = {
    "WILMER PACHECO HIRAOKA": "#1645df",
    "MIGUEL VARGAS CASTRO": "#f56b18",
    "RONALD GAMARRA SANTILLAN": "#00897b",
}


def color_supervisor(nombre):
    return COLORES_SUPERVISOR.get(str(nombre).strip().upper(), "#999999")


def leyenda_supervisores(fig, datos):
    for nombre in sorted(datos.Supervisor.dropna().unique()):
        fig.add_trace(go.Scatter(x=[None], y=[None], mode="markers",
            marker=dict(color=color_supervisor(nombre), size=9), name=nombre,
            legendgroup=nombre, hoverinfo="skip"))


def numero(v):
    try:
        x = float(str(v).strip().replace(",", "."))
        return x if math.isfinite(x) else math.nan
    except (TypeError, ValueError):
        return math.nan


def fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def recuperacion(f, c, t):
    # Balance de dos productos; todas las leyes deben estar en g/tm.
    if not all(math.isfinite(x) for x in (f, c, t)):
        return math.nan
    if f <= 0 or not (c > f >= t >= 0):
        return math.nan
    return 100 * c * (f - t) / (f * (c - t))


def leer_excel(archivo, hoja=HOJA):
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=False)
    try:
        if hoja not in wb.sheetnames:
            raise ValueError(f"No existe {hoja}. Hojas: {wb.sheetnames}")
        filas = list(wb[hoja].iter_rows(min_row=1, max_row=30, values_only=True))
        def normal(v):
            return str(v or "").strip().upper().replace(" ", "")
        candidatos = [(i, sum(fecha(v) is not None for v in row)) for i, row in enumerate(filas)]
        fila_fecha, cantidad = max(candidatos, key=lambda x: x[1])
        if cantidad < 2:
            raise ValueError("No se encontró una fila con fechas completas de Excel.")
        indices = {}
        for nombre, clave in (("Faja", "FAJA"), ("OVER", "OVER"),
                              ("BULK", "CONC.BULK"), ("Relave", "RVE.FINAL"),
                              ("Rec", "RECUP(%)")):
            indices[nombre] = [i for i, row in enumerate(filas) if len(row) > 2 and normal(row[2]) == clave]
            if nombre != "Rec" and len(indices[nombre]) != 2:
                raise ValueError(f"Se esperaban dos filas de {nombre}, día y noche.")
        # Expandir solo los bloques combinados reales de supervisores.
        supervisores = {}
        for bloque in wb[hoja].merged_cells.ranges:
            valor = wb[hoja].cell(bloque.min_row, bloque.min_col).value
            if isinstance(valor, str) and regex.match(r"GUARDIA\s*\(", valor.strip(), regex.I):
                for columna in range(bloque.min_col, bloque.max_col + 1):
                    supervisores[(bloque.min_row - 1, columna - 1)] = valor.strip()
        for i, row in enumerate(filas):
            for j, valor in enumerate(row):
                if isinstance(valor, str) and regex.match(r"GUARDIA\s*\(", valor.strip(), regex.I):
                    supervisores[(i, j)] = valor.strip()
        filas_supervisor = sorted({i for i, _ in supervisores})
        registros, vistos = [], set()
        for col, v in enumerate(filas[fila_fecha]):
            f = fecha(v)
            if f is None:
                continue
            unidad = normal(filas[fila_fecha + 2][col])
            prox = normal(filas[fila_fecha + 2][col+1]) if col+1 < len(filas[fila_fecha+2]) else ""
            au = normal(filas[fila_fecha+1][col])
            if not (unidad in ("G/TM", "GR/TM") or (not unidad and au == "AU" and prox == "OZ/TM")):
                continue
            if f in vistos:
                raise ValueError(f"Fecha duplicada: {f}")
            vistos.add(f)
            for turno, guardia in enumerate(("Día", "Noche")):
                anteriores = [i for i in filas_supervisor if i < indices["Faja"][turno]]
                fila_supervisor = max(anteriores) if anteriores else -1
                texto = supervisores.get((fila_supervisor, col), "")
                coincidencia = regex.match(r"GUARDIA\s*\(([^)]+)\)\s*(.*)", texto, regex.I)
                equipo = coincidencia.group(1).strip() if coincidencia else "Sin asignación"
                supervisor = coincidencia.group(2).strip() if coincidencia else "Sin asignación"
                valores = {n: numero(filas[indices[n][turno]][col]) for n in ("Faja", "OVER", "BULK", "Relave")}
                cg = valores["BULK"]
                onzas = numero(filas[indices["BULK"][turno]][col+1]) if col+1 < len(filas[0]) else math.nan
                # Las fórmulas de onzas pueden devolver cero cuando g/tm está vacío.
                if not math.isfinite(cg):
                    onzas = math.nan
                elif not math.isfinite(onzas):
                    onzas = cg / 31.10348013
                re = numero(filas[indices["Rec"][turno]][col]) if len(indices["Rec"]) == 2 else math.nan
                rc = recuperacion(valores["OVER"], cg, valores["Relave"])
                valido_excel = math.isfinite(re) and 0 <= re <= 100 and math.isfinite(rc)
                r = re if valido_excel else rc
                origen = "Excel" if valido_excel else ("Calculada con leyes" if math.isfinite(rc) else "Sin datos / revisar leyes")
                registros.append({"Fecha": f, "Guardia": guardia,
                    "Equipo": equipo, "Supervisor": supervisor,
                    "Faja (g/tm)": valores["Faja"], "OVER (g/tm)": valores["OVER"],
                    "Relave (g/tm)": valores["Relave"], "BULK (g/tm)": cg,
                    "BULK (Oz/tm)": onzas, "Recuperación (%)": r,
                    "Rec. Excel (%)": re, "Rec. calculada (%)": rc, "Origen recuperación": origen})
    finally:
        wb.close()
    if not registros:
        raise ValueError("No se encontraron resultados con fechas y unidades reconocibles.")
    df = pd.DataFrame(registros)
    df["Fecha"] = pd.to_datetime(df["Fecha"])
    return df.loc[df.Fecha >= FECHA_INICIO].sort_values(["Fecha", "Guardia"]).reset_index(drop=True)


def formato(v, unidad):
    return "Sin datos" if pd.isna(v) else f"{v:.3f} {unidad}"


def estilo(fig, fechas, titulo, unidad):
    fig.update_layout(title=titulo, template="plotly_white", height=620,
        xaxis_title="Fecha", yaxis_title=unidad,
        legend=dict(orientation="h", y=-.32), margin=dict(b=165))
    fig.update_xaxes(tickmode="array", tickvals=fechas,
        ticktext=[pd.Timestamp(f).strftime("%d/%m") for f in fechas], tickangle=-90)
    fig.update_yaxes(tickformat=".3f", rangemode="tozero")
    return fig


def eje_supervisores(fig, datos, fechas, guardias):
    """Un rótulo por fecha; distingue responsables de día y noche."""
    etiquetas = []
    for f in fechas:
        partes = [pd.Timestamp(f).strftime("%d/%m/%Y")]
        for guardia in guardias:
            sub = datos[(datos.Fecha == f) & (datos.Guardia == guardia)]
            nombre = sub.Supervisor.iloc[0] if not sub.empty else None
            equipo = sub.Equipo.iloc[0] if not sub.empty else None
            if pd.notna(nombre):
                partes.append(f"{guardia} · {equipo}: {nombre}")
        etiquetas.append("<br>".join(partes))
    fig.update_xaxes(tickmode="array", tickvals=fechas, ticktext=etiquetas,
                     tickangle=-90, tickfont=dict(size=10), automargin=True,
                     title_text="Fecha · turno · supervisor")
    fig.update_layout(height=950, margin=dict(b=390),
                      legend=dict(
                          orientation="h",
                          y=-0.50,
                          yanchor="top",
                          x=0
                          ))
    return fig


def grafica_leyes(datos, fechas, guardias, ref_relave, ref_bulk, etiquetas):
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    colores = {("Relave", "Día"): "#1645df", ("Relave", "Noche"): "#f56b18",
               ("OVER", "Día"): "#00897b", ("OVER", "Noche"): "#8c6d31",
               ("BULK", "Día"): "#8e24aa", ("BULK", "Noche"): "#c2185b"}
    for producto, unidad in (("Relave", "g/tm"), ("OVER", "g/tm"), ("BULK", "Oz/tm")):
        col = f"{producto} ({unidad})"
        for guardia in guardias:
            sub = datos[datos.Guardia == guardia].sort_values("Fecha")
            mostrar = etiquetas == "Todas" or (etiquetas == "Solo relave" and producto == "Relave")
            if producto in ("Relave", "BULK", "OVER"):
                # Colorear cada tramo por el responsable de su fecha inicial.
                # Se mantienen espacios cuando faltan resultados o se filtra un supervisor.
                for i in range(len(sub) - 1):
                    p1, p2 = sub.iloc[i], sub.iloc[i + 1]
                    if pd.isna(p1[col]) or pd.isna(p2[col]):
                        continue
                    fig.add_trace(go.Scatter(
                        x=[p1.Fecha, p2.Fecha], y=[p1[col], p2[col]],
                        mode="lines",
                        line=dict(color=color_supervisor(p1.Supervisor), width=2,
                                  dash="dash" if producto == "OVER" else "solid"),
                        showlegend=False, hoverinfo="skip"), secondary_y=producto == "BULK")
            fig.add_trace(go.Scatter(x=sub.Fecha, y=sub[col],
                name=f"{producto} {guardia} · {unidad}",
                mode=(("markers+text" if mostrar else "markers") if producto in ("Relave", "BULK", "OVER")
                      else ("lines+markers+text" if mostrar else "lines+markers")),
                text=[f"{v:.3f}" if pd.notna(v) else "" for v in sub[col]] if mostrar else None,
                textposition="top center" if guardia == "Día" else "bottom center",
                textfont=dict(size=10, color=[color_supervisor(v) for v in sub.Supervisor]),
                marker=dict(color=[color_supervisor(v) for v in sub.Supervisor], size=8,
                    symbol=("circle" if producto == "Relave" else "triangle-up" if producto == "OVER" else "square") + ("-open" if guardia == "Noche" else "")),
                line=dict(color=colores[(producto, guardia)], width=2, dash="dash" if producto == "OVER" else "solid"),
                connectgaps=False,
                customdata=pd.DataFrame({"Equipo": sub.Equipo.fillna("Sin datos"),
                    "Supervisor": sub.Supervisor.fillna("Sin datos"),
                    "RECUP": sub["Recuperación (%)"].map(lambda v: f"{v:.3f}" if pd.notna(v) else "Sin datos")}).to_numpy(),
                hovertemplate=f"%{{x|%d/%m/%Y}}<br>%{{y:.3f}} {unidad}<br>Equipo %{{customdata[0]}} · %{{customdata[1]}}<br>RECUP: %{{customdata[2]}} %<extra>%{{fullData.name}}</extra>"),
                secondary_y=producto == "BULK")
    fig.add_hline(y=ref_relave, line_dash="dash", line_color="green", secondary_y=False,
        annotation_text=f"Ref. relave {ref_relave:.3f} g/tm", annotation_position="bottom left")
    fig.add_hline(y=ref_bulk, line_dash="dashdot", line_color="#8e24aa", secondary_y=True,
        annotation_text=f"Ref. BULK {ref_bulk:.3f} Oz/tm", annotation_position="top right")
    estilo(fig, fechas, "Leyes de oro: Relave, OVER y CONC. BULK", "Relave y OVER (g/tm)")
    fig.update_yaxes(title_text="CONC. BULK (Oz/tm)", secondary_y=True, showgrid=False)
    eje_supervisores(fig, datos, fechas, guardias)
    leyenda_supervisores(fig, datos)
    return fig


def grafica_barras(datos, fechas, guardias, variable, etiquetas=True):
    fig = go.Figure()
    for guardia in guardias:
        sub = datos[datos.Guardia == guardia].sort_values("Fecha")
        fig.add_trace(go.Bar(x=sub.Fecha, y=sub[variable], name=guardia,
            marker_color=[color_supervisor(v) for v in sub.Supervisor],
            marker_pattern_shape="/" if guardia == "Noche" else "",
            text=[f"{v:.3f}" if pd.notna(v) else "" for v in sub[variable]] if etiquetas else None,
            textposition="outside", cliponaxis=False,
            customdata=sub[["Equipo", "Supervisor"]].fillna("Sin asignación").to_numpy(),
            hovertemplate="%{x|%d/%m/%Y}<br>%{y:.3f}<br>Equipo %{customdata[0]} · %{customdata[1]}<extra>%{fullData.name}</extra>"))
    estilo(fig, fechas, f"{variable} por fecha y guardia", variable)
    fig.update_layout(barmode="group")
    if variable == "Recuperación (%)":
        fig.update_yaxes(range=[0, 110])
    eje_supervisores(fig, datos, fechas, guardias)
    leyenda_supervisores(fig, datos)
    return fig


def main():
    st.set_page_config(page_title="Control de flotación", page_icon="📊", layout="wide")
    st.title("📊 Control de flotación — Leyes y recuperación de Au")
    st.markdown("**Made by Ing. Arturo De la Cruz Coaquira**")
    st.markdown(
        "[jdelacruz@unsa.edu.pe](mailto:jdelacruz@unsa.edu.pe) · "
        "[jorge.delacruz.c@uni.pe](mailto:jorge.delacruz.c@uni.pe) · "
        "Teléfono: [958897458](tel:+51958897458)"
    )
    candidatos = sorted(p for p in CARPETA.glob("*.xls*") if not p.name.startswith("~$") and p.suffix.lower() in (".xlsx", ".xlsm"))
    if not candidatos:
        st.error("Coloque el Excel .xlsm o .xlsx junto a este código.")
        st.stop()
    nombres = [p.name for p in candidatos]
    nombre = st.sidebar.selectbox("Excel", nombres, index=nombres.index(NOMBRE_EXCEL) if NOMBRE_EXCEL in nombres else 0)
    archivo = CARPETA / nombre
    st.sidebar.button("🔄 Actualizar datos")
    try:
        df = leer_excel(archivo)
    except Exception as e:
        st.error(f"No se pudo leer el archivo: {e}")
        st.stop()
    con_datos = df.dropna(subset=["Relave (g/tm)", "OVER (g/tm)", "BULK (g/tm)"], how="all")
    if con_datos.empty:
        st.warning("No hay resultados de leyes.")
        st.stop()
    ultima = con_datos.Fecha.max()
    st.caption(f"Excel: {archivo.resolve()} · Último guardado: {datetime.fromtimestamp(archivo.stat().st_mtime):%d/%m/%Y %H:%M} · Último dato: {ultima:%d/%m/%Y}")
    modo = st.sidebar.radio("Período", ["Rango de fechas", "Mes"])
    if modo == "Mes":
        anios = sorted(con_datos.Fecha.dt.year.unique().tolist())
        anio = st.sidebar.selectbox("Año", anios, index=len(anios)-1)
        meses = sorted(con_datos.loc[con_datos.Fecha.dt.year == anio, "Fecha"].dt.month.unique().tolist())
        mes = st.sidebar.selectbox("Mes", meses, index=len(meses)-1, format_func=lambda x: MESES[x])
        inicio = pd.Timestamp(anio, mes, 1)
        fin = inicio + pd.offsets.MonthEnd(0)
    else:
        rango = st.sidebar.date_input("Desde / hasta", value=(con_datos.Fecha.min().date(), ultima.date()),
            min_value=con_datos.Fecha.min().date(), max_value=ultima.date())
        if len(rango) != 2:
            st.info("Seleccione inicio y fin.")
            st.stop()
        inicio, fin = map(pd.Timestamp, rango)
    guardias = st.sidebar.multiselect("Guardias", ["Día", "Noche"], default=["Día", "Noche"])
    if not guardias:
        st.info("Seleccione al menos una guardia.")
        st.stop()
    rr = st.sidebar.number_input("Referencia relave (g/tm)", min_value=0., value=.65, step=.05, format="%.3f")
    rb = st.sidebar.number_input("Referencia BULK (Oz/tm)", min_value=0., value=3.3, step=.1, format="%.3f")
    etiquetas = st.sidebar.selectbox("Etiquetas de curvas", ["Todas", "Solo relave", "Ninguna"])
    puntos = df[(df.Fecha >= inicio) & (df.Fecha <= fin) & df.Guardia.isin(guardias)].copy()
    supervisores_disponibles = sorted(puntos.Supervisor.unique().tolist())
    elegidos = st.sidebar.multiselect("Supervisores", supervisores_disponibles, default=supervisores_disponibles)
    puntos = puntos[puntos.Supervisor.isin(elegidos)].copy()
    # Gráficas con fechas completas; faltantes como espacios, nunca como cero.
    fechas = pd.date_range(inicio, fin)
    indice = pd.MultiIndex.from_product([fechas, guardias], names=["Fecha", "Guardia"])
    datos = puntos.set_index(["Fecha", "Guardia"]).reindex(indice).reset_index()
    st.subheader(f"{inicio:%d/%m/%Y} — {fin:%d/%m/%Y}")
    st.caption("Wilmer = azul · Miguel = naranja · Ronald = verde. Todas las líneas toman el color del supervisor de la fecha inicial de cada tramo. BULK y relave: continuos; OVER: discontinuo. Círculo = relave, triángulo = OVER, cuadrado = BULK; marcador abierto y barras rayadas = noche. Valores con tres decimales.")
    st.dataframe(puntos[["Equipo", "Supervisor"]].drop_duplicates().sort_values("Equipo"), hide_index=True, width="stretch")
    for guardia in guardias:
        sub = puntos[puntos.Guardia == guardia]
        columnas = st.columns(4)
        columnas[0].metric(f"Relave {guardia}", formato(sub["Relave (g/tm)"].mean(), "g/tm"))
        columnas[1].metric(f"OVER {guardia}", formato(sub["OVER (g/tm)"].mean(), "g/tm"))
        columnas[2].metric(f"BULK {guardia}", formato(sub["BULK (Oz/tm)"].mean(), "Oz/tm"))
        columnas[3].metric(f"Recuperación media {guardia}", formato(sub["Recuperación (%)"].mean(), "%"))
    st.caption("Promedios aritméticos de resultados disponibles; no son balances mensuales ponderados por tonelaje o metal.")
    with st.expander("Cómo se obtiene la recuperación", expanded=False):
        st.latex(r"R_{Au}(\%)=100\frac{c(f-t)}{f(c-t)}")
        st.write("f = OVER, c = CONC. BULK y t = relave, todas las leyes en g/tm. Es un balance de dos productos del circuito alimentado con OVER; no es la recuperación global desde FAJA.")
        st.write("Se usa la recuperación del Excel si es válida y existen las tres leyes. Si falta la fórmula o su valor guardado, se calcula con las leyes. La tabla identifica el origen. Se requieren c > f ≥ t ≥ 0 y f > 0; los casos incompatibles quedan vacíos para revisión.")
    tabs = st.tabs(["Leyes · dos ejes", "Barras por guardia", "Recuperación (%)", "Comparación mensual", "Datos y revisión"])
    with tabs[0]:
        st.plotly_chart(grafica_leyes(datos, fechas, guardias, rr, rb, etiquetas), width="stretch")
    with tabs[1]:
        variable = st.selectbox("Variable de barras", ["Relave (g/tm)", "OVER (g/tm)", "BULK (Oz/tm)", "Recuperación (%)"])
        fig = grafica_barras(datos, fechas, guardias, variable)
        if variable == "Relave (g/tm)":
            fig.add_hline(y=rr, line_color="green", line_dash="dash", annotation_text=f"Ref. {rr:.3f} g/tm")
        elif variable == "BULK (Oz/tm)":
            fig.add_hline(y=rb, line_color="#8e24aa", line_dash="dash", annotation_text=f"Ref. {rb:.3f} Oz/tm")
        st.plotly_chart(fig, width="stretch")
    with tabs[2]:
        st.subheader("RECUP (%) por fecha, turno y supervisor")
        st.plotly_chart(grafica_barras(datos, fechas, guardias, "Recuperación (%)"), width="stretch")
        st.dataframe(puntos.groupby(["Equipo", "Supervisor", "Guardia"])["Recuperación (%)"].agg(["count", "mean", "min", "max"]).round(3), width="stretch")
        st.dataframe(puntos.groupby("Guardia")["Recuperación (%)"].agg(["count", "mean", "min", "max"]).round(3), width="stretch")
    with tabs[3]:
        variable_mes = st.selectbox("Comparar promedios mensuales", ["Relave (g/tm)", "OVER (g/tm)", "BULK (Oz/tm)", "Recuperación (%)"])
        anual = df[(df.Fecha.dt.year == inicio.year) & df.Guardia.isin(guardias) & df.Supervisor.isin(elegidos)].copy()
        anual["Mes"] = anual.Fecha.dt.month
        resumen = anual.groupby(["Mes", "Supervisor", "Guardia"])[variable_mes].mean().unstack(["Supervisor", "Guardia"]).dropna(how="all")
        fig = go.Figure()
        for supervisor, guardia in resumen.columns:
            valores = resumen[(supervisor, guardia)]
            fig.add_trace(go.Bar(x=[MESES[m] for m in resumen.index], y=valores,
                name=f"{supervisor} · {guardia}", marker_color=color_supervisor(supervisor),
                marker_pattern_shape="/" if guardia == "Noche" else "",
                text=valores.round(3), textposition="auto"))
        fig.update_layout(title=f"{variable_mes} — promedios de {inicio.year}", barmode="group", template="plotly_white", yaxis_title=variable_mes)
        st.plotly_chart(fig, width="stretch")
        st.dataframe(resumen.rename(index=lambda x: MESES[x]).round(3), width="stretch")
    with tabs[4]:
        tabla = puntos.copy()
        tabla["Fecha"] = tabla.Fecha.dt.strftime("%d/%m/%Y")
        st.dataframe(tabla.round(3), width="stretch", hide_index=True)
        st.download_button("Descargar CSV del período", tabla.to_csv(index=False).encode("utf-8-sig"),
            file_name=f"Flotacion_{inicio:%Y%m%d}_{fin:%Y%m%d}.csv", mime="text/csv")
        discrepancia = (puntos["Rec. Excel (%)"] - puntos["Rec. calculada (%)"]).abs() > .1
        if discrepancia.any():
            st.warning("Hay diferencias mayores a 0.1 puntos porcentuales entre la recuperación guardada y la recalculada. Recalcule y guarde el Excel.")
            st.dataframe(puntos.loc[discrepancia], hide_index=True)
        incompletos = puntos[["OVER (g/tm)", "BULK (g/tm)", "Relave (g/tm)"]].notna().any(axis=1) & puntos["Recuperación (%)"].isna()
        if incompletos.any():
            st.warning(f"{incompletos.sum()} guardias con leyes incompletas o incompatibles: recuperación sin calcular.")
    st.sidebar.caption("Guarde el Excel y pulse Actualizar datos. No se usa caché.")


if __name__ == "__main__":
    main()
